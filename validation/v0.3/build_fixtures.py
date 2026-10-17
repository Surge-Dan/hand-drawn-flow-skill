"""Original synthetic fixtures; not a production universal parser."""
from pathlib import Path
from zipfile import ZipFile, ZIP_STORED, ZIP_DEFLATED
import json
import xml.etree.ElementTree as ET
from docx import Document
from docx.shared import Pt
from reportlab.pdfgen import canvas
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.cidfonts import UnicodeCIDFont
from PIL import Image, ImageDraw, ImageFont
from pypdf import PdfReader

ROOT = Path(__file__).parent
SRC = ROOT / 'fixtures'
SRC.mkdir(exist_ok=True)
doc = Document()
doc.styles['Normal'].font.name = 'Microsoft YaHei'
doc.styles['Normal'].font.size = Pt(11)
doc.add_heading('共享工作室预约需求', 0)
doc.add_paragraph('原创虚构测试材料。用于检验并行审核与异常出口，不代表真实产品规则。')
doc.add_heading('预约流程', 1)
doc.add_paragraph('用户提交预约后，空间管理员与设备管理员并行审核。两者都通过，才进入待支付。任一拒绝则结束本次申请。')
t = doc.add_table(rows=1, cols=2)
t.rows[0].cells[0].text = '条件'
t.rows[0].cells[1].text = '结果'
for a,b in [('两者都通过','待支付'),('任一拒绝','已拒绝'),('30分钟内支付','预约成功'),('超时未支付','自动取消')]:
    c=t.add_row().cells;c[0].text=a;c[1].text=b
doc.add_paragraph('取消后释放时段；不自动重新申请。')
doc.save(SRC/'booking.docx')

pdfmetrics.registerFont(UnicodeCIDFont('STSong-Light'))
c = canvas.Canvas(str(SRC/'design.pdf'), pagesize=(595,842))
c.setFont('STSong-Light',18);c.drawString(48,790,'预约服务设计说明')
c.setFont('STSong-Light',11)
c.drawString(48,762,'原创虚构测试材料。概念结构示意，非真实部署。')
for x,txt in [(48,'客户端'),(225,'预约服务'),(420,'时段库')]:
    c.rect(x,640,110,55);c.drawString(x+14,662,txt)
for x1,x2 in [(158,225),(335,420)]:
    c.line(x1,667,x2,667);c.line(x2,667,x2-7,671);c.line(x2,667,x2-7,663)
c.drawString(48,610,'箭头表示请求方向；客户端不得直接写时段库。')
rows=[('状态','时段处理'),('待支付','锁定30分钟'),('预约成功','保留占用'),('自动取消','释放时段')]
for i,(a,b) in enumerate(rows):
    y=560-i*38;c.rect(48,y-12,440,38);c.line(220,y-12,220,y+26)
    c.drawString(60,y,a);c.drawString(235,y,b)
c.save()

chapters=[('z-intro','范围','本书为原创虚构测试短篇，讨论怎样把收藏的资料用于实际问题。不提供效果数据。'),
('a-save','收藏','保存链接只是保留入口，不代表已经理解。收藏后还需要明确想回答的问题。'),
('m-question','问题','先写一个具体问题，再寻找与该问题相关的资料。例子：怎样减少会议后的遗漏？'),
('b-evidence','证据','将相关片段与来源一起留下，比较相同点、差异和适用条件。无关资料暂留在收藏区，不塞进结论。'),
('x-action','试用','用证据写出一个小行动，在自己的情境中试用，记录结果，再回到问题检查。试用结果不自动证明普遍有效。'),
('c-limit','限制','这是一种工作方法，不保证效率提升。图可以概括问题、证据、试用和回看，不能把收藏量画成能力增长。')]
with ZipFile(SRC/'reading.epub','w') as z:
    z.writestr('mimetype','application/epub+zip',compress_type=ZIP_STORED)
    z.writestr('META-INF/container.xml','<container xmlns="urn:oasis:names:tc:opendocument:xmlns:container" version="1.0"><rootfiles><rootfile full-path="OEBPS/content.opf" media-type="application/oebps-package+xml"/></rootfiles></container>')
    manifest=''.join(f'<item id="{i}" href="{i}.xhtml" media-type="application/xhtml+xml"/>' for i,_,_ in chapters)
    spine=''.join(f'<itemref idref="{i}"/>' for i,_,_ in chapters)
    z.writestr('OEBPS/content.opf',f'<package xmlns="http://www.idpf.org/2007/opf" version="3.0" unique-identifier="book"><metadata xmlns:dc="http://purl.org/dc/elements/1.1/"><dc:identifier id="book">synthetic-reading-v1</dc:identifier><dc:title>从收藏到试用</dc:title><dc:language>zh</dc:language></metadata><manifest>{manifest}</manifest><spine>{spine}</spine></package>')
    for i,title,body in chapters:
        z.writestr(f'OEBPS/{i}.xhtml',f'<html xmlns="http://www.w3.org/1999/xhtml"><head><title>{title}</title></head><body><h1>{title}</h1><p>{body}</p></body></html>',compress_type=ZIP_DEFLATED)

font=ImageFont.truetype('C:/Windows/Fonts/msyh.ttc',30)
cards=[('card-a.png','卡片 C｜次日清晨',['小岚把修好的蓝伞放回门边。','伞柄系着白纸条：修好了。']),('card-z.png','卡片 A｜傍晚',['小岚在门边发现一把蓝伞。','一根伞骨断了，她把伞带进屋。']),('card-m.png','卡片 B｜当晚',['小岚在桌前接好断掉的伞骨。','窗外仍在下雨。'])]
for name,title,lines in cards:
    im=Image.new('RGB',(1000,430),'white');d=ImageDraw.Draw(im)
    for y,s in zip([35,115,185,320],[title,*lines,'原创虚构故事卡；以卡片时间为序']):d.text((35,y),s,font=font,fill='black')
    im.save(SRC/name)

# Read from saved files, not the creation variables.
extracted={}
read_doc=Document(SRC/'booking.docx')
ns={'w':'http://schemas.openxmlformats.org/wordprocessingml/2006/main'}
extracted['docx']=[{'block':i+1,'kind':e.tag.split('}')[-1], 'text':' | '.join(e.xpath('.//w:t/text()'))} for i,e in enumerate(read_doc.element.body) if e.tag.split('}')[-1] in ('p','tbl')]
extracted['pdf']=[{'file_page':i+1,'text':p.extract_text()} for i,p in enumerate(PdfReader(SRC/'design.pdf').pages)]
with ZipFile(SRC/'reading.epub') as z:
    container=ET.fromstring(z.read('META-INF/container.xml'))
    opfpath=next(e.attrib['full-path'] for e in container.iter() if e.tag.endswith('rootfile'))
    opf=ET.fromstring(z.read(opfpath));base=Path(opfpath).parent
    manifest={e.attrib['id']:e.attrib['href'] for e in opf.iter() if e.tag.endswith('}item')}
    extracted['epub']=[]
    for e in opf.iter():
        if e.tag.endswith('}itemref'):
            path=(base/manifest[e.attrib['idref']]).as_posix();body=ET.fromstring(z.read(path))
            extracted['epub'].append({'spine_id':e.attrib['idref'],'path':path,'text':' '.join(body.itertext())})
assert [x['spine_id'] for x in extracted['epub']]==['z-intro','a-save','m-question','b-evidence','x-action','c-limit']
assert any('两者都通过' in x['text'] for x in extracted['docx'])
assert '客户端不得直接写时段库' in extracted['pdf'][0]['text']
(ROOT/'extracted.json').write_text(json.dumps(extracted,ensure_ascii=False,indent=2),encoding='utf-8')
print('Created four input types. Three saved-file extraction assertions passed. Images require separate visual inspection.')
